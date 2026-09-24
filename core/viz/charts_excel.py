"""Native Excel chart export (legacy-parity, FR-6.5).

The legacy NLP Suite exported charts as real Excel workbooks
(``charts_Excel_util.py``): a "Data" worksheet holding the series values and
categories, a first-positioned "Chart" worksheet holding a native openpyxl
chart (bar / line / pie / scatter / radar / bubble), axis titles, category
labels pinned at the bottom (``tickLblPos="low"``) and the legend removed for
a single series. This module reproduces that output from a prepared chart
(:class:`core.viz.chartspec.PreparedChart`), so charts opened in Excel look
the way the original NLP Suite's Excel charts look.

Deliberately NOT ported: the macro-enabled ``.xlsm`` hover-over path. It
required shipping binary ``.xlsm`` templates with embedded VBA (a macro-
security liability); the HTML (Plotly) charts cover hover needs interactively.
The Excel export is the static native-chart view with identical visuals.

openpyxl supports the kinds the legacy GUI offered: bar, line, pie, scatter,
radar, bubble. The remaining kinds (histogram/box/heatmap/sunburst/treemap/
violin/waffle/calendar) have no native Excel chart type; requesting Excel
export for them fails with a precise diagnostic (``CHART_EXCEL_UNSUPPORTED``)
instead of silently drawing something else.
"""

from __future__ import annotations

import io
import math
from typing import TYPE_CHECKING, Any

from core.result import Diagnostic, Result

if TYPE_CHECKING:
    from openpyxl.chart import Series as OpenpyxlSeries

    from core.viz.chartspec import ChartSpec, PreparedChart

__all__ = ["EXCEL_KINDS", "excel_chart_bytes"]

# Excel chart types openpyxl can write natively (the legacy GUI's menu).
EXCEL_KINDS: tuple[str, ...] = ("bar", "line", "pie", "scatter", "radar", "bubble")

_EXCEL_FIX = 'pip install "nlp-suite-ng[excel]"'
_SINGLE_SERIES = 2


def excel_chart_bytes(prepared: PreparedChart, spec: ChartSpec) -> Result[bytes]:
    """Render the prepared chart as a native Excel workbook (xlsx bytes).

    Workbook layout mirrors the legacy output: sheet "Chart" first (so Excel
    opens on the chart), sheet "Data" behind it with categories in column 1
    and one value column per series. A grouped chart contributes one series
    per group, exactly as the legacy multi-series workbook.
    """
    if spec.kind not in EXCEL_KINDS:
        return Result.failure(_unsupported_excel_kind(spec))
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        return Result.failure(
            Diagnostic.error(
                "CHART_EXCEL_OPENPYXL_MISSING",
                f"Excel export needs openpyxl: {exc}",
                fix=_EXCEL_FIX,
            )
        )
    frame = prepared.data
    if frame.empty:
        return Result.failure(Diagnostic.error("CHART_EMPTY", "frame is empty"))

    try:
        workbook = Workbook()
        data_sheet = workbook.active
        data_sheet.title = "Data"
        n_series = _write_data_sheet(data_sheet, frame, spec)
        chart = _build_chart(data_sheet, frame, prepared, spec, n_series)
        if chart is None:
            return Result.failure(Diagnostic.error("CHART_EXCEL_FAILED", "Excel chart could not be built"))
        chart_sheet = workbook.create_sheet("Chart")
        chart_sheet.add_chart(chart, "A1")
        # Excel opens on the first sheet; the legacy moved Chart to the front.
        workbook.move_sheet("Chart", offset=-1)
        buffer = io.BytesIO()
        workbook.save(buffer)
    except Exception as exc:  # openpyxl raises broadly on odd cell values
        return Result.failure(Diagnostic.error("CHART_EXCEL_FAILED", f"Excel chart failed: {exc}"))
    return Result.success(buffer.getvalue())


def _unsupported_excel_kind(spec: ChartSpec) -> Diagnostic:
    kind = spec.kind
    if kind in ("sunburst", "treemap", "violin", "waffle", "calendar"):
        return Diagnostic.error(
            "CHART_EXCEL_UNSUPPORTED",
            f"Excel has no native {kind} chart. Use --format html (interactive) or --format png "
            "(static image); the Excel export covers bar/line/pie/scatter/radar/bubble.",
        )
    if kind in ("histogram", "box"):
        return Diagnostic.error(
            "CHART_EXCEL_UNSUPPORTED",
            f"Excel has no native {kind} chart object (the legacy suite exported these "
            "Plotly-only too). Use --format html, or --kind bar for an Excel bar chart.",
        )
    if kind == "heatmap":
        return Diagnostic.error(
            "CHART_EXCEL_UNSUPPORTED",
            "Excel heatmaps are conditional formatting, not chart objects. Use --format html "
            "(interactive heatmap) or work from chart_data.csv directly.",
        )
    return Diagnostic.error("CHART_EXCEL_UNSUPPORTED", f"no Excel chart for kind {kind!r}")


# ---------------------------------------------------------------------------
# Data sheet: categories in column 1, one value column per series
# ---------------------------------------------------------------------------


def _write_data_sheet(sheet: Any, frame: Any, spec: ChartSpec) -> int:
    """Fill the Data sheet; return the number of value series written."""
    if spec.kind == "pie" or "group" not in frame.columns or frame["group"].nunique() <= 1:
        sheet.append([str(spec.x), str(spec.y)])
        for row in frame.itertuples(index=False):
            sheet.append([_cell_value(row.x), _cell_value(row.y)])
        return 1
    groups = sorted(frame["group"].astype(str).unique().tolist())
    sheet.append([str(spec.x), *groups])
    keys = list(dict.fromkeys(frame["x"].tolist()))
    lookup = {(row.x, str(row.group)): row.y for row in frame.itertuples(index=False)}
    for key in keys:
        sheet.append([_cell_value(key), *(_cell_value(lookup.get((key, group))) for group in groups)])
    return len(groups)


def _cell_value(value: Any) -> Any:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if hasattr(value, "isoformat"):  # datetime: write as ISO text (tz-safe)
        return str(value)
    return value


# ---------------------------------------------------------------------------
# Chart construction
# ---------------------------------------------------------------------------


def _build_chart(data_sheet: Any, frame: Any, prepared: PreparedChart, spec: ChartSpec, n_series: int) -> Any:
    """Create the openpyxl chart object against the filled data sheet."""
    from openpyxl.chart import BarChart, BubbleChart, LineChart, PieChart, RadarChart, ScatterChart

    builders: dict[str, tuple[Any, bool]] = {
        "bar": (BarChart(), True),
        "line": (LineChart(), True),
        "pie": (PieChart(), False),
        "scatter": (ScatterChart(), True),
        "radar": (RadarChart(), False),
        "bubble": (BubbleChart(), True),
    }
    chart, with_axes = builders[spec.kind]
    if spec.kind == "radar":
        chart.type = "marker"
    if spec.kind == "bar":
        chart.type = "bar" if spec.horizontal else "col"
        if spec.bar_mode == "stack":
            chart.grouping = "stacked"
            chart.overlap = 100
        elif spec.bar_mode == "relative":
            chart.grouping = "percentStacked"
            chart.overlap = 100

    from openpyxl.chart import Reference, Series

    rows = sum(1 for _ in frame.itertuples(index=False))
    if n_series <= 1:
        # Categories in col 1, values in col 2 (legacy add_data/set_categories shape).
        data = Reference(data_sheet, min_col=2, min_row=1, max_row=1 + rows)
        categories = Reference(data_sheet, min_col=1, min_row=2, max_row=1 + rows)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
    else:
        groups = sorted(frame["group"].astype(str).unique().tolist())
        for index in range(n_series):
            values = Reference(data_sheet, min_col=2 + index, min_row=1, max_row=1 + rows)
            chart.series.append(Series(values, title=groups[index]))
        categories = Reference(data_sheet, min_col=1, min_row=2, max_row=1 + rows)
        chart.set_categories(categories)

    _cosmetics(chart, prepared, spec, with_axes=with_axes)
    return chart


def _series_count(frame: Any, spec: ChartSpec) -> int:
    if spec.kind == "pie" or "group" not in frame.columns:
        return 1
    return int(frame["group"].nunique()) or 1


def _cosmetics(chart: Any, prepared: PreparedChart, spec: ChartSpec, *, with_axes: bool) -> None:
    """Shared legacy cosmetology: title, axis titles, labels low, 1-series legend off."""
    chart.title = spec.title or (f"{spec.y} by {spec.x}")
    if with_axes:
        x_title = spec.x_label or str(spec.x)
        y_title = spec.y_label or str(prepared.y_label or spec.y)
        if x_title:
            # The legacy appended blank lines so long category labels had
            # room below the axis (its insertLines behavior).
            chart.x_axis.title = x_title + "\n\n"
            chart.x_axis.tickLblPos = "low"
            chart.x_axis.tickLblSkip = 1
        if y_title:
            chart.y_axis.title = y_title
    if len(chart.series) < _SINGLE_SERIES:
        # The legacy hid the legend for single-series charts ("Series 1" box).
        chart.legend = None
