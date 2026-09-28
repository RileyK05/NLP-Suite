"""Document details from a spreadsheet the user already has (docs/internal/PLAN_0.5.0.md 1.4.2).

A researcher often keeps a table beside the texts -- one row per speech, with
its party, its occasion, its chapter number. This module reads that table and
says, before anything is stored, which row belongs to which document and what
each column would become. Pure: bytes in, a report out. The desktop's
``fields/import`` route stores the values.

Matching is deliberately forgiving and deliberately loud. A cell matches a
document by exact name, then by the name without its extension, then ignoring
case; every row that matched nothing and every document no row named is
reported, never silently dropped.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
import csv
from dataclasses import dataclass, field
from datetime import date, datetime
import io
from pathlib import PurePath
import re

import pandas as pd

from core.result import Diagnostic, Result

__all__ = [
    "PREVIEW_ROWS",
    "ColumnPlan",
    "MatchReport",
    "match_metadata",
    "proposed_target",
    "read_table",
]

#: How many matched rows the preview shows.
PREVIEW_ROWS = 10
#: A column of names and at least one of details.
_MIN_COLUMNS = 2
#: How many unmatched names a warning lists.
_LISTED = 5
#: Tables bigger than this are not a list of documents' details.
MAX_TABLE_BYTES = 10 * 1024 * 1024

# Header words that name the axis. Proposed, never applied silently: the
# preview shows each mapping with a menu so a "number" column that counts
# something else can be kept as itself.
_DATE_HEADERS = frozenset({"date", "year", "published", "when", "day"})
_ORDER_HEADERS = frozenset({"order", "chapter", "number", "index", "no", "position", "sequence", "#"})
_NAME_HEADERS = ("document", "file", "filename", "file name", "name", "text")
_NUMBER = re.compile(r"^-?\d+(?:\.\d+)?$")
_ISO_DAY = re.compile(r"^\d{4}(?:-\d{2}(?:-\d{2})?)?$")


@dataclass(frozen=True, slots=True)
class ColumnPlan:
    """What one spreadsheet column would become.

    ``target`` is the detail's name ("Date", "Order", "Party"); empty means
    the column is left out. ``kind`` is what its values look like (``date``,
    ``number`` or ``text``), so the preview can say why a mapping was proposed.
    """

    column: str
    target: str
    kind: str
    filled: int
    samples: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class MatchReport:
    """Everything the preview shows, and the values an apply would store."""

    name_column: str
    #: "" when rows name documents by file name; else the detail they are keyed by ("Speaker").
    matched_on: str
    columns: tuple[ColumnPlan, ...]
    #: document name -> {detail name: value}, empty cells left out.
    values: dict[str, dict[str, str]]
    #: How each matched document was found: "exact", "without extension", "ignoring case", "by Speaker".
    matched_by: dict[str, str]
    unmatched_rows: tuple[str, ...]
    unmatched_documents: tuple[str, ...]
    #: Rows naming a document an earlier row already named (the first row wins).
    duplicate_rows: tuple[str, ...]
    #: The first matched rows as they would be stored, for the preview table.
    preview: tuple[tuple[str, dict[str, str]], ...] = field(default=())


def read_table(data: bytes, filename: str) -> Result[pd.DataFrame]:
    """A CSV, TSV or XLSX file as a table of text cells.

    Every cell is read as text: "01" must stay a chapter called "01", and a
    year column must not become 1934.0. XLSX needs openpyxl, which the suite
    installs for its Excel charts; without it the file is refused with a note
    to save it as CSV.
    """
    if len(data) > MAX_TABLE_BYTES:
        return Result.failure(
            Diagnostic.error("METADATA_TOO_LARGE", f"That file is over {MAX_TABLE_BYTES // (1024 * 1024)} MB.")
        )
    suffix = PurePath(filename).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return _read_workbook(data)
    if suffix not in (".csv", ".tsv", ".txt", ""):
        return Result.failure(
            Diagnostic.error("METADATA_BAD_FORMAT", f"Choose a .csv, .tsv or .xlsx file, not {suffix or filename}.")
        )
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
    if not text.strip():
        return Result.failure(Diagnostic.error("METADATA_EMPTY", "That file is empty."))
    sample = text[:4096]
    try:
        delimiter = "\t" if suffix == ".tsv" else csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = ","
    try:
        frame = pd.read_csv(io.StringIO(text), sep=delimiter, dtype=str, keep_default_na=False)
    except (pd.errors.ParserError, ValueError) as exc:
        return Result.failure(Diagnostic.error("METADATA_UNREADABLE", f"That file could not be read as a table: {exc}"))
    return _tidy(frame)


def _read_workbook(data: bytes) -> Result[pd.DataFrame]:
    try:
        from openpyxl import load_workbook  # noqa: PLC0415 - the Excel component is optional
    except ImportError:
        return Result.failure(
            Diagnostic.error(
                "METADATA_NEEDS_OPENPYXL",
                "Reading .xlsx needs the Excel component (Setup page). Or save the sheet as CSV and choose that.",
            )
        )
    try:
        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        rows = [[_cell_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
    except Exception as exc:  # openpyxl raises broadly on damaged files
        return Result.failure(Diagnostic.error("METADATA_UNREADABLE", f"That workbook could not be read: {exc}"))
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        return Result.failure(Diagnostic.error("METADATA_EMPTY", "That workbook's first sheet is empty."))
    width = max(len(row) for row in rows)
    header = [cell or f"Column {index + 1}" for index, cell in enumerate(rows[0] + [""] * (width - len(rows[0])))]
    body = [row + [""] * (width - len(row)) for row in rows[1:]]
    return _tidy(pd.DataFrame(body, columns=header))


def _cell_text(value: object) -> str:
    """A workbook cell as the text a reader sees: dates as ISO days, 3.0 as 3."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _tidy(frame: pd.DataFrame) -> Result[pd.DataFrame]:
    frame = frame.rename(columns=lambda column: str(column).strip())
    frame = frame.loc[:, [column for column in frame.columns if column and not column.startswith("Unnamed:")]]
    frame = frame.fillna("").astype(str).apply(lambda column: column.str.strip())
    frame = frame.loc[(frame != "").any(axis=1)].reset_index(drop=True)
    if frame.empty or len(frame.columns) < _MIN_COLUMNS:
        return Result.failure(
            Diagnostic.error(
                "METADATA_TOO_NARROW",
                "The table needs a column of document names and at least one column of details.",
            )
        )
    return Result.success(frame)


def proposed_target(column: str) -> str:
    """The detail a column header suggests: "Year" -> "Date", "Chapter" -> "Order", "party" -> "Party"."""
    key = column.strip().casefold()
    if key in _DATE_HEADERS:
        return "Date"
    if key in _ORDER_HEADERS:
        return "Order"
    cleaned = re.sub(r"[^A-Za-z0-9 _-]", " ", column).strip()
    cleaned = re.sub(r"\s+", " ", cleaned)[:40].strip()
    return cleaned[:1].upper() + cleaned[1:] if cleaned else ""


def _kind(values: Sequence[str]) -> str:
    filled = [value for value in values if value]
    if filled and all(_ISO_DAY.match(value) for value in filled):
        return "date"
    if filled and all(_NUMBER.match(value) for value in filled):
        return "number"
    return "text"


def _stem(name: str) -> str:
    return PurePath(name).stem if PurePath(name).suffix else name


class _Matcher:
    """Finds the documents a cell names.

    By name: exact, then without extension, then ignoring case -- one
    document. By a detail (``key``): every document whose detail has that
    value, ignoring case -- a row keyed by "Harry S Truman" gives Party to
    all of his speeches.
    """

    def __init__(self, names: Sequence[str], details: Mapping[str, Mapping[str, str]], key: str = "") -> None:
        self.key = key
        self.exact = {name: name for name in names}
        self.stem: dict[str, str] = {}
        self.folded: dict[str, str] = {}
        for name in names:
            self.stem.setdefault(_stem(name), name)
            self.folded.setdefault(_stem(name).casefold(), name)
        self.by_value: dict[str, list[str]] = {}
        if key:
            for name in names:
                value = next((v for k, v in details.get(name, {}).items() if k.casefold() == key.casefold()), "")
                if value:
                    self.by_value.setdefault(value.strip().casefold(), []).append(name)

    def find(self, cell: str) -> list[tuple[str, str]]:
        if not cell:
            return []
        if self.key:
            return [(name, f"by {self.key}") for name in self.by_value.get(cell.strip().casefold(), [])]
        if cell in self.exact:
            return [(self.exact[cell], "exact")]
        stem = _stem(cell)
        if stem in self.stem:
            return [(self.stem[stem], "without extension")]
        if stem.casefold() in self.folded:
            return [(self.folded[stem.casefold()], "ignoring case")]
        return []

    def covered(self, cells: Sequence[str]) -> int:
        return len({name for cell in cells for name, _how in self.find(cell)})


def _best_key(frame: pd.DataFrame, matchers: Sequence[_Matcher], name_column: str | None) -> tuple[str, _Matcher, int]:
    """The column and the way of matching that reach the most documents.

    Names win a tie over a detail, and a "file"/"name" header wins a tie over
    another column, so an ordinary sheet of file names is read as one.
    """
    columns = [name_column] if name_column is not None else list(frame.columns)

    def reach(column: str, matcher: _Matcher) -> int:
        cells = [str(cell) for cell in frame[column]]
        filled = [cell.strip().casefold() for cell in cells if cell.strip()]
        # A key names one thing per row. Through a detail, a column whose
        # values repeat ("Democratic", "Democratic") is a value, not a key:
        # keyed on it, every Democratic speech would take the first
        # Democratic row's president.
        if matcher.key and len(set(filled)) < len(filled):
            return 0
        return matcher.covered(cells)

    scored = [
        (
            reach(column, matcher),
            0 if matcher.key else 1,
            1 if column.strip().casefold() in _NAME_HEADERS else 0,
            index,
        )
        for index, (column, matcher) in enumerate((c, m) for c in columns for m in matchers)
    ]
    best = max(scored)
    column, matcher = [(c, m) for c in columns for m in matchers][best[3]]
    return column, matcher, best[0]


def match_metadata(
    frame: pd.DataFrame,
    names: Sequence[str],
    *,
    details: Mapping[str, Mapping[str, str]] | None = None,
    name_column: str | None = None,
    targets: Mapping[str, str] | None = None,
) -> Result[MatchReport]:
    """Which spreadsheet row belongs to which documents, and what each column becomes.

    *names* are the project's document names; *details* is each document's
    existing details (name -> {detail: value}), so a sheet keyed by Speaker
    rather than by file name also matches. *name_column* picks the key
    column; left out, the column (and the way of matching) that reaches the
    most documents is used. *targets* maps a column to its detail name (""
    leaves it out); columns it does not mention get :func:`proposed_target`.
    """
    if name_column is not None and name_column not in frame.columns:
        return Result.failure(
            Diagnostic.error(
                "METADATA_NO_COLUMN", f"The table has no column {name_column!r}; it has {list(frame.columns)}."
            )
        )
    known = details or {}
    keys = sorted({key for mine in known.values() for key in mine}, key=str.casefold)
    matchers = [_Matcher(names, known)] + [_Matcher(names, known, key) for key in keys]
    chosen, matcher, reached = _best_key(frame, matchers, name_column)
    if reached == 0:
        examples = ", ".join(repr(cell) for cell in frame.iloc[:3, 0])
        return Result.failure(
            Diagnostic.error(
                "METADATA_NO_MATCH",
                f"No column names this project's documents (the first column starts {examples}; the documents "
                f"are named like {names[0]!r}). One column must hold the file names, with or without .txt, or "
                "the values of a detail the documents already have, such as Speaker."
                if names
                else "This project has no documents to match.",
            )
        )
    plans = _plans(frame, chosen, targets or {})
    kept = [plan for plan in plans if plan.target]
    seen_targets: dict[str, str] = {}
    for plan in kept:
        other = seen_targets.setdefault(plan.target.casefold(), plan.column)
        if other != plan.column:
            return Result.failure(
                Diagnostic.error(
                    "METADATA_SAME_TARGET",
                    f"Columns {other!r} and {plan.column!r} would both become {plan.target!r}; leave one out or "
                    "rename it.",
                )
            )

    values: dict[str, dict[str, str]] = {}
    matched_by: dict[str, str] = {}
    unmatched_rows: list[str] = []
    duplicates: list[str] = []
    for _, row in frame.iterrows():
        cell = str(row[chosen])
        found = matcher.find(cell)
        if not found:
            if cell:
                unmatched_rows.append(cell)
            continue
        if all(document in values for document, _how in found):
            duplicates.append(cell)
            continue
        for document, how in found:
            if document not in values:
                values[document] = {
                    plan.target: _value(str(row[plan.column])) for plan in kept if str(row[plan.column])
                }
                matched_by[document] = how

    notes = _notes(kept, unmatched_rows, duplicates)
    missing = tuple(name for name in names if name not in values)
    preview = tuple((name, values[name]) for name in names if name in values)[:PREVIEW_ROWS]
    return Result.success(
        MatchReport(
            name_column=chosen,
            matched_on=matcher.key,
            columns=tuple(plans),
            values=values,
            matched_by=matched_by,
            unmatched_rows=tuple(unmatched_rows),
            unmatched_documents=missing,
            duplicate_rows=tuple(duplicates),
            preview=preview,
        ),
        *notes,
    )


def _plans(frame: pd.DataFrame, key_column: str, targets: Mapping[str, str]) -> list[ColumnPlan]:
    """What every column but the key would become."""
    plans: list[ColumnPlan] = []
    for column in frame.columns:
        if column == key_column:
            continue
        cells = [str(value) for value in frame[column]]
        plans.append(
            ColumnPlan(
                column=column,
                target=targets[column].strip() if column in targets else proposed_target(column),
                kind=_kind(cells),
                filled=sum(1 for cell in cells if cell),
                samples=tuple(dict.fromkeys(cell for cell in cells if cell))[:_LISTED],
            )
        )
    return plans


def _notes(kept: Sequence[ColumnPlan], unmatched_rows: Sequence[str], duplicates: Sequence[str]) -> list[Diagnostic]:
    """What the preview must say before anything is stored."""
    notes = [
        Diagnostic.warning(
            "METADATA_DATE_FORMAT",
            f"Column {plan.column!r} would become the Date, but not every value is written as 1934, 1934-01 or "
            f"1934-01-03 (for example {next((s for s in plan.samples if not _ISO_DAY.match(s)), '')!r}); those "
            "documents would stay undated.",
        )
        for plan in kept
        if plan.target.casefold() == "date" and plan.kind != "date"
    ]
    if unmatched_rows:
        notes.append(
            Diagnostic.warning(
                "METADATA_UNMATCHED_ROWS",
                f"{len(unmatched_rows)} row(s) name no document in this project and will be left out: "
                + ", ".join(unmatched_rows[:_LISTED])
                + ("…" if len(unmatched_rows) > _LISTED else ""),
            )
        )
    if duplicates:
        notes.append(
            Diagnostic.warning(
                "METADATA_DUPLICATE_ROWS",
                f"{len(duplicates)} row(s) name documents earlier rows already named; the first row is used.",
            )
        )
    return notes


def _value(raw: str) -> str:
    """One cell as the detail stores it: a year for a Date column stays "1934", "3.0" becomes "3"."""
    text = raw.strip()
    if _NUMBER.match(text) and "." in text and float(text).is_integer():
        return str(int(float(text)))
    return text
